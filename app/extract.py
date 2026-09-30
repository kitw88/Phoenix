import io
from pathlib import Path


def extract_text(filename: str, raw: bytes) -> str:
    suffix = Path(filename or "").suffix.lower()
    if suffix in {".txt", ".md", ".json"}:
        text = raw.decode("utf-8", errors="replace")
    elif suffix == ".docx":
        text = _docx_text(raw)
    elif suffix == ".pdf":
        text = _pdf_text(raw)
    elif suffix in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".tif", ".tiff", ".bmp"}:
        raise ValueError("圖片裡的文字這一步抽不出來。請改放有文字層的 pdf、docx 或文字檔。")
    else:
        raise ValueError("只接受文字、docx 或 pdf。")
    if not text.strip():
        raise ValueError("這份檔沒有抽出文字。")
    return text.strip()


def _docx_text(raw: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(raw))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(part for part in parts if part and part.strip())


def _pdf_text(raw: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(raw))
    return "\n".join((page.extract_text() or "") for page in reader.pages)
