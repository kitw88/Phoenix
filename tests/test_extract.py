import io

from docx import Document

from app.extract import extract_text
from app.postman_run import split_chat_body


def test_plain_text():
    assert extract_text("sheet.txt", "FCN coupon".encode()) == "FCN coupon"


def test_docx_includes_table_cells():
    document = Document()
    document.add_paragraph("Term sheet")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Product"
    table.rows[0].cells[1].text = "FCN"
    buffer = io.BytesIO()
    document.save(buffer)
    text = extract_text("sheet.docx", buffer.getvalue())
    assert "Term sheet" in text
    assert "FCN" in text


def test_image_is_rejected():
    try:
        extract_text("page.png", b"not-an-image")
    except ValueError as exc:
        assert "圖片" in str(exc)
    else:
        raise AssertionError("expected ValueError")


def test_split_chat_body_keeps_reasoning_apart_from_json():
    body = """
    {
      "choices": [
        {
          "message": {
            "content": "{\\"product\\":\\"FCN\\",\\"matched_features\\":[\\"coupon\\"]}",
            "reasoning_content": "coupon barrier"
          }
        }
      ]
    }
    """
    output, reasoning = split_chat_body(body)
    assert '"product":"FCN"' in output
    assert reasoning == "coupon barrier"
